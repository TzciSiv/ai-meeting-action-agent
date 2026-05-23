import unittest
from io import BytesIO

from docx import Document

from backend.transcript_import import extract_docx_transcript


def make_docx_bytes(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=max(len(row) for row in table_rows))
        for row_index, row in enumerate(table_rows):
            for cell_index, value in enumerate(row):
                table.cell(row_index, cell_index).text = value

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class DocxTranscriptImportTests(unittest.TestCase):
    def test_extracts_paragraphs_and_table_rows(self):
        content = make_docx_bytes(
            ["Alice: We need the launch checklist.", "Ben: I will send it today."],
            [["Owner", "Task"], ["Ben", "Send checklist"]],
        )

        transcript = extract_docx_transcript(BytesIO(content), "meeting.docx")

        self.assertIn("Alice: We need the launch checklist.", transcript)
        self.assertIn("Ben: I will send it today.", transcript)
        self.assertIn("Ben | Send checklist", transcript)

    def test_rejects_non_docx_files(self):
        with self.assertRaisesRegex(ValueError, "docx"):
            extract_docx_transcript(BytesIO(b"plain text"), "meeting.txt")

    def test_rejects_docx_without_text(self):
        content = make_docx_bytes([])

        with self.assertRaisesRegex(ValueError, "readable transcript text"):
            extract_docx_transcript(BytesIO(content), "empty.docx")


if __name__ == "__main__":
    unittest.main()
