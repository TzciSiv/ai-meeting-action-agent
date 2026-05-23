from io import BytesIO
from typing import BinaryIO

from docx import Document


MAX_DOCX_BYTES = 15 * 1024 * 1024


def _document_text(document: Document) -> str:
    lines: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    # Some meeting templates store transcript rows in tables, so keep those
    # rows readable instead of silently dropping them.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines).strip()


def extract_docx_transcript(uploaded_file: BinaryIO, filename: str) -> str:
    if not filename.lower().endswith(".docx"):
        raise ValueError("Please upload a .docx transcript file.")

    content = uploaded_file.read()
    if not content:
        raise ValueError("The DOCX file is empty.")
    if len(content) > MAX_DOCX_BYTES:
        raise ValueError("The DOCX file is too large. Please upload a file under 15 MB.")

    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise ValueError("Could not read this DOCX file.") from exc

    transcript = _document_text(document)
    if not transcript:
        raise ValueError("The DOCX file did not contain readable transcript text.")

    return transcript
