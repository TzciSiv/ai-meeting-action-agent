import re
from io import BytesIO
from typing import Any

from docx import Document


GENERIC_EVIDENCE = "structured action item returned by GPT action extraction."


def clean_markdown_text(text: str) -> str:
    text = re.sub(r"^#+\s*", "", text.strip())
    text = re.sub(r"^[-*]\s*", "", text)
    return text.strip()


def add_markdown_to_doc(document: Document, markdown_text: str) -> None:
    for raw_line in markdown_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("## "):
            document.add_heading(clean_markdown_text(line), level=2)
        elif line.startswith("- "):
            document.add_paragraph(clean_markdown_text(line), style="List Bullet")
        else:
            document.add_paragraph(clean_markdown_text(line))


def document_bytes(document: Document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_transcript_docx(transcript: str) -> bytes:
    document = Document()
    document.add_heading("Meeting Transcript", level=1)
    document.add_paragraph(transcript.strip() or "No transcript available.")
    return document_bytes(document)


def build_agent_report_docx(
    transcript: str,
    summary_markdown: str,
    actions: list[dict[str, Any]],
    decisions: list[str],
    risks: list[str],
    follow_up_question: str = "",
    follow_up_answer: str = "",
) -> bytes:
    document = Document()
    document.add_heading("AI Meeting Action Agent Report", level=1)

    document.add_heading("Action Board", level=2)
    if not actions:
        document.add_paragraph("No action items were found.")
    for index, item in enumerate(actions, start=1):
        document.add_paragraph(f"{index}. {item.get('task', 'Unspecified task')}", style="List Number")
        document.add_paragraph(f"Owner: {item.get('owner', 'Unassigned')}")
        document.add_paragraph(f"Deadline: {item.get('deadline', 'Not mentioned')}")
        document.add_paragraph(f"Status: {item.get('status', 'Open')}")
        item_evidence = str(item.get("evidence", "")).strip()
        if item_evidence and item_evidence.lower() != GENERIC_EVIDENCE:
            document.add_paragraph(f"Evidence: {item_evidence}")

    document.add_heading("Decisions", level=2)
    for decision in decisions or ["None captured"]:
        document.add_paragraph(str(decision), style="List Bullet")

    document.add_heading("Risks", level=2)
    for risk in risks or ["None detected"]:
        document.add_paragraph(str(risk), style="List Bullet")

    if follow_up_question or follow_up_answer:
        document.add_heading("Follow-up Answer", level=2)
        if follow_up_question:
            document.add_paragraph(f"Question: {follow_up_question}")
        document.add_paragraph(follow_up_answer or "Not asked")

    document.add_heading("Meeting Summary", level=2)
    add_markdown_to_doc(document, summary_markdown)

    document.add_heading("Transcript", level=2)
    document.add_paragraph(transcript.strip() or "No transcript available.")
    return document_bytes(document)
